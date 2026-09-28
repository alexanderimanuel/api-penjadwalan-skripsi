"""Read-back audit and distributable archive; never changes timetable content."""
import argparse
import csv
import io
import json
import math
from pathlib import Path
import zipfile

from export_august import read_json
from final_experiment import atomic_json, digest, load_frozen
from model_august import ROOT


def main():
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('folder');args=parser.parse_args()
    folder=Path(args.folder).resolve()
    model,_,_,_=load_frozen(ROOT/'results/final-pilot/school-user-approved-v1/FROZEN_EXPERIMENT_CONFIG.json')
    load_frozen(ROOT/'results/final-pilot/technical-v1/FROZEN_EXPERIMENT_CONFIG.json')
    for p in folder.rglob('*.csv'):
        # Canonicalize only record separators, not values or timetable placement.
        old=p.read_bytes();new=old.replace(b'\r\r\n',b'\r\n')
        if new!=old:
            parse=lambda value:list(csv.DictReader(io.StringIO(value.decode('utf-8'),newline='')))
            assert parse(old)==parse(new)
            p.write_bytes(new)
    # Read only: xlsx authoring remains solely in artifact-tool.
    import openpyxl
    workbook=openpyxl.load_workbook(folder/'Review-SMAN8.xlsx',read_only=False,data_only=True)
    facts=read_json(folder/'workbook-data.json')
    comparison=workbook['Perbandingan']
    assert len(workbook.sheetnames)==12
    for col,label in enumerate(['BASELINE','A','B','C'],2):
        assert comparison.cell(5,col).value==facts['metrics'][label]['H']
        assert math.isclose(comparison.cell(13,col).value,facts['metrics'][label]['S'],rel_tol=1e-12)
    assert [comparison.cell(35,c).value for c in range(3,6)]==['N/A']*3
    review=workbook['Isian Sekolah']
    assert review['B16'].value is None
    assert all(review.cell(r,c).value is None for r in range(8,12) for c in range(2,6))
    assert len(review.data_validations.dataValidation)>=2
    workbook.close()
    # Every source teacher name must be absent from all distributed text and XLSX XML.
    ignored={'workbook-data.json','PACKAGE_AUDIT.json'}
    files=[p for p in folder.rglob('*') if p.is_file() and 'verification' not in p.relative_to(folder).parts
           and p.name not in ignored and not p.name.endswith('.inspect.ndjson')]
    text=''.join(p.read_text(encoding='utf-8') for p in files if p.suffix in ('.json','.csv','.html','.md'))
    with zipfile.ZipFile(folder/'Review-SMAN8.xlsx') as z:
        text+=''.join(z.read(name).decode('utf-8') for name in z.namelist() if name.endswith('.xml'))
    for t in model.data['teachers']:
        if t['display_name'] in text:raise ValueError('Teacher name leaked into review package')
    for label in ['BASELINE','A','B','C']:
        rows=read_json(folder/label/'timetable.json')
        assert len(rows)==513 and len({r['event_id'] for r in rows})==513
        assert all('source_metadata' not in r for r in rows)
    review_input=read_json(folder/'school-review-input.json')
    assert all(v is None for k,v in review_input.items() if k!='ratings')
    assert all(v is None for r in review_input['ratings'] for k,v in r.items() if k!='schedule')
    audit={'status':'PASS','workbook_sheets':12,'schedules':4,'class_views':108,'teacher_views':232,
           'source_teacher_names_absent':True,'reviewer_answers_blank':True,'xlsx_readback_values_and_validation':'PASS',
           'raw_and_seed_reproduction':read_json(folder/'reproduction.json'),
           'previous_frozen_integrity':'PASS','file_hashes':{p.relative_to(folder).as_posix():digest(p) for p in files}}
    atomic_json(folder/'PACKAGE_AUDIT.json',audit)
    archive=folder.parent/'Paket-Peninjauan-SMAN8-Malang-v1.zip'
    with zipfile.ZipFile(archive,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=9) as z:
        for p in [*files,folder/'PACKAGE_AUDIT.json']:z.write(p,p.relative_to(folder).as_posix())
    print(json.dumps({'status':'PASS','archive':str(archive),'files':len(files)+1,'bytes':archive.stat().st_size},indent=2))


if __name__=='__main__':main()
